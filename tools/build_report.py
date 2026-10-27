"""Build the private submission report from actual captured screen images."""
import json
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

root = Path(__file__).resolve().parents[1]
screens = root / 'artifacts' / 'screenshots'
output = root.parent / 'submission' / 'QE190017.docx'
credentials = json.loads((root / 'test-account.local.json').read_text(encoding='utf-8'))
items = [
    ('firebase-auth', 'Firebase console — Authentication', 'The Users list includes the dedicated grader account.'),
    ('firebase-data', 'Firebase console — Firestore data', 'The books collection contains five sample books belonging to the grader account.'),
    ('firebase-rules', 'Firebase console — Security Rules', 'The published rules restrict access by owner and validate canonical book IDs and field values.'),
    ('02-register-validation', 'Register screen — validation', 'Submitting the empty registration form shows the required-field messages.'),
    ('03-register-success', 'Register screen — success', 'A real Firebase account was created and Registration successful is displayed.'),
    ('05-login', 'Login screen', 'The empty sign-in form.'),
    ('06-list-after-login', 'List screen — after login', 'The signed-in account, Logout action and saved books are shown.'),
    ('07-add-validation', 'Add Book form — validation', 'Submitting an empty book form shows the required-field messages.'),
    ('08-add-success', 'Add Book form — success', 'Book added is shown after the new book is saved.'),
    ('14-genre-filter', 'Filters and search — genre filter', 'Selecting Fiction limits the visible books to that genre.'),
    ('15-status-filter', 'Filters and search — status filter', 'Reading combines with the genre filter.'),
    ('16-favorites-filter', 'Filters and search — favorites only', 'Only matching favorite books remain visible.'),
    ('17-search', 'Filters and search — search', 'The title search combines with all active filters and ignores surrounding spaces.'),
    ('18-summary', 'Summary', 'Books and Favorites counts and the status breakdown describe the books currently shown.'),
    ('09-detail-view', 'Detail screen — view mode', 'Title, author, genre, status, rating and notes are displayed.'),
    ('10-detail-edit-validation', 'Detail screen — edit mode validation', 'Saving an empty title displays Title is required.'),
    ('11-detail-after-save', 'Detail screen — after Save', 'The revised notes are saved and the screen returns to view mode.'),
    ('12-favorite-before', 'Favorite toggle — before', 'The verification book is not yet marked as a favorite.'),
    ('13-favorite-after', 'Favorite toggle — after', 'The filled star indicates the favorite value saved in Firestore.'),
    ('19-delete-confirmation', 'Delete — confirmation', 'The confirmation dialog provides Cancel and Delete actions.'),
    ('20-after-delete', 'Delete — after', 'The verification book is removed and the original five books remain.'),
    ('04-empty-state', 'Empty state', 'A fresh second account displays No books found.'),
]
missing = [name for name, _, _ in items if not (screens / f'{name}.png').exists()]
if missing:
    raise SystemExit('Missing actual screenshots: ' + ', '.join(missing))

doc = Document()
section = doc.sections[0]
section.top_margin = section.bottom_margin = Inches(.6)
section.left_margin = section.right_margin = Inches(.65)
section.page_width, section.page_height = Inches(8.27), Inches(11.69)
for name in ['Normal', 'Title', 'Heading 1', 'Heading 2']:
    style = doc.styles[name]
    style.font.name = 'Arial'
    style.font.size = Pt(11 if name == 'Normal' else 16)
    style.font.color.rgb = RGBColor(0, 0, 0)
doc.styles['Normal'].paragraph_format.space_after = Pt(10)

# Keep the exact field names and order required by the exam template.
for line in [
    'Student Code: QE190017',
    'Full Name: Nguyễn Công Nhất',
    'GitHub Repository URL: https://github.com/nguyenhat3105/bookshelf-qe190017',
    'APK Download Link: https://github.com/nguyenhat3105/bookshelf-qe190017/releases/download/v1.0.0/app-release.apk',
    'Firebase Project ID: bookshelf-qe190017',
    f'Test account email: {credentials["email"]}',
    f'Test account password: {credentials["password"]}',
]:
    doc.add_paragraph(line)

for name, caption, explanation in items:
    doc.add_page_break()
    doc.add_heading(caption, level=1)
    doc.add_paragraph(explanation)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(screens / f'{name}.png'), width=Inches(6.85 if name.startswith('firebase-') else 3.45))

output.parent.mkdir(parents=True, exist_ok=True)
doc.save(output)
print(f'Report created with {len(items)} actual screenshots: {output}')
